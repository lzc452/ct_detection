# -*- coding: utf-8 -*-
"""
Web 端 #3（失败态）+ #6（文件台账）+ #8（阈值直传）集成测试。

不依赖真实模型：注入一个假 ultralytics 模块 + 临时 models 目录，然后用 Flask test_client
打真实路由。验证的是"接口契约"，不是检测精度：

1. 批量 3 张（1 有缺陷 / 1 无缺陷 / 1 损坏），返回 results 长度恒为 3 且 index 与输入一一对应
   —— 旧实现会把损坏那张从数组里删掉，导致缩略图与文件名错位；
2. 状态与计数：ok_defect=1 / ok_clean=1 / failed=1，失败文件名可枚举；
3. 文件台账 file_records 每个输入文件恰好一条，Excel 汇总的"检测文件数"= 3（旧口径恒为 1）；
4. 阈值直传：模型收到的 conf 就是用户值（0.157），不再被 max(0.25, x) 吞掉；
5. parse_conf 对 nan/inf/abc/越界值的兜底。

用法：runtime\\python.exe scripts/test_web_ledger.py
"""
import importlib.util
import io
import json
import hashlib
import os
import shutil
import sys
import tempfile
import types
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]


# ==================== 假 ultralytics ====================
class FakeTensor:
    """模拟 torch.Tensor 的最小接口：.cpu().numpy().astype(...) / .tolist()"""

    def __init__(self, arr):
        import numpy as np
        self.a = np.asarray(arr)

    def cpu(self):
        return self

    def numpy(self):
        return self.a

    def astype(self, t):
        return FakeTensor(self.a.astype(t))

    def tolist(self):
        return self.a.tolist()

    def __len__(self):
        return len(self.a)


class FakeBoxes:
    def __init__(self, n):
        import numpy as np
        self.xyxy = FakeTensor(np.array([[10, 10, 30, 30]] * n, dtype=float) if n else np.zeros((0, 4)))
        self.cls = FakeTensor(np.array([1] * n, dtype=float) if n else np.zeros((0,)))
        self.conf = FakeTensor(np.array([0.9] * n, dtype=float) if n else np.zeros((0,)))

    def __len__(self):
        return len(self.xyxy)

    def __getitem__(self, mask):
        return FakeBoxes(int(mask.sum()) if hasattr(mask, 'sum') else 0)


class FakeResult:
    """亮度 > 10 的图判为 1 个缺陷框，纯黑图判为无缺陷（够用来跑通状态机）"""

    def __init__(self, img):
        import numpy as np
        arr = np.asarray(img)
        self._img = arr
        n = 1 if (arr.size and float(arr.mean()) > 10) else 0
        self.boxes = FakeBoxes(n)

    def plot(self):
        return self._img.copy()


class FakeYOLO:
    calls = []

    def __init__(self, path, task=None):
        self.path = path

    def __call__(self, source, **kw):
        FakeYOLO.calls.append(dict(kw))
        if isinstance(source, (list, tuple)):
            return [FakeResult(x) for x in source]
        return [FakeResult(source)]

    def predict(self, source, **kw):
        return self(source, **kw)


def install_fake_ultralytics():
    m = types.ModuleType('ultralytics')
    m.YOLO = FakeYOLO
    sys.modules['ultralytics'] = m
    # ultralytics 的部分子模块可能被间接导入，一并兜底
    for sub in ('ultralytics.nn', 'ultralytics.utils'):
        sys.modules.setdefault(sub, types.ModuleType(sub))


def jpeg_bytes(rgb_value, size=(64, 64)):
    from PIL import Image
    buf = io.BytesIO()
    Image.new('RGB', size, (rgb_value, rgb_value, rgb_value)).save(buf, 'JPEG')
    buf.seek(0)
    return buf


def main():
    tmp = Path(tempfile.mkdtemp(prefix='web_ledger_'))
    model_dir = tmp / 'models'
    results_dir = tmp / 'results'
    model_dir.mkdir()
    results_dir.mkdir()
    (model_dir / 'fake_v9.pt').write_bytes(b'fake')
    (model_dir / '_train_meta.json').write_text(json.dumps(dict(
        hyperparams=dict(imgsz=1920), roi=False, roi_k=0.0,
        best_pt_sha256=hashlib.sha256(b'fake').hexdigest())), encoding='utf8')

    os.environ['CT_MODEL_DIR'] = str(model_dir)
    os.environ['CT_RESULTS_DIR'] = str(results_dir)
    os.environ['CT_NO_BROWSER'] = '1'
    os.environ.setdefault('PROCESSOR_ARCHITECTURE', 'AMD64')

    install_fake_ultralytics()
    spec = importlib.util.spec_from_file_location('ct_web_app', str(PROJECT_ROOT / 'web' / 'app.py'))
    app_module = importlib.util.module_from_spec(spec)
    sys.modules['ct_web_app'] = app_module
    spec.loader.exec_module(app_module)
    client = app_module.app.test_client()

    fails = []

    def check(cond, msg):
        if not cond:
            fails.append(msg)
            print('  [FAIL] ' + msg)

    try:
        # ---------- 1) parse_conf（#8） ----------
        default = app_module.CONF_DEFAULT
        cases = [('nan', default), ('inf', default), ('abc', default), (None, default),
                 ('0.157', 0.157), ('1.5', 1.0), ('-0.3', 0.0)]
        for raw, want in cases:
            got = app_module.parse_conf(raw)
            check(abs(got - want) < 1e-9, 'parse_conf(%r) = %r，应为 %r' % (raw, got, want))

        # ---------- 2) 批量：1 缺陷 + 1 干净 + 1 损坏 ----------
        # 把损坏文件放在**中间**：旧实现直接 continue 跳过，后面的干净图结果会前移一位，
        # 前端按 chunk[k] 取缩略图就张冠李戴——这是 #3 最危险的次生 bug
        names = ['B/正面/0001.jpg', 'B/正面/broken.jpg', 'B/正面/0002.jpg']
        files = [
            (jpeg_bytes(200), '0001.jpg'),
            (io.BytesIO(b'this is not an image'), 'broken.jpg'),
            (jpeg_bytes(0), '0002.jpg'),
        ]
        data = {'conf': '0.157', 'batch_id': 'BT1',
                'file_paths': json.dumps(names), 'files': files}
        r = client.post('/api/detect/batch', data=data, content_type='multipart/form-data')
        resp = r.get_json()
        check(r.status_code == 200 and resp.get('success'), '批量接口失败：%s' % resp)

        res = resp.get('results', [])
        check(len(res) == 3, 'results 长度应为 3（含失败），实际 %d' % len(res))
        check([x.get('index') for x in res] == [0, 1, 2], 'results 下标不对齐：%s' % [x.get('index') for x in res])
        check([x.get('filename') for x in res] == names, 'results 文件名与输入不一致：%s' % [x.get('filename') for x in res])
        stats = {x['index']: x.get('status') for x in res}
        check(stats.get(0) == 'ok_defect', '第 1 张应 ok_defect，实际 %s' % stats.get(0))
        check(stats.get(1) == 'failed', '第 2 张（损坏）应 failed，实际 %s' % stats.get(1))
        check(stats.get(2) == 'ok_clean', '第 3 张应 ok_clean（不能被前移错位），实际 %s' % stats.get(2))
        check(resp.get('ok_defect') == 1 and resp.get('ok_clean') == 1 and resp.get('failed') == 1,
              '计数不对：%s' % {k: resp.get(k) for k in ('ok_defect', 'ok_clean', 'failed')})
        check(resp.get('failed_files') == ['B/正面/broken.jpg'], '失败清单不对：%s' % resp.get('failed_files'))

        # ---------- 3) 阈值直传 ----------
        last = FakeYOLO.calls[-1] if FakeYOLO.calls else {}
        check(abs(last.get('conf', -1) - 0.157) < 1e-9,
              '模型没有收到用户阈值，实收 conf=%r（旧实现恒为 0.25）' % last.get('conf'))

        # ---------- 4) 台账 ----------
        recs = [f for f in app_module.file_records if f.get('批次ID') == 'BT1']
        check(len(recs) == 3, '文件台账应为 3 条，实际 %d' % len(recs))
        by_name = {f['文件路径']: f for f in recs}
        check(by_name[names[0]]['状态'] == 'ok_defect' and by_name[names[0]]['缺陷数'] == 1,
              '台账：%s' % by_name.get(names[0]))
        check(by_name[names[1]]['状态'] == 'failed' and by_name[names[1]]['备注'],
              '台账没有记录失败原因：%s' % by_name.get(names[1]))
        check(all(abs(f['阈值'] - 0.157) < 1e-9 for f in recs), '台账阈值没记对')

        # ---------- 5) Excel 汇总 ----------
        r = client.post('/api/export', json={'batch_id': 'BT1'})
        exp = r.get_json()
        check(exp.get('success'), '导出失败：%s' % exp)
        xlsx = results_dir / Path(exp['download_url']).name
        check(xlsx.exists(), '报表文件没生成')
        from openpyxl import load_workbook
        wb = load_workbook(xlsx)
        check('文件台账' in wb.sheetnames, '缺少"文件台账"页：%s' % wb.sheetnames)
        ws = wb['统计汇总']
        summary = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value
                   for i in range(1, ws.max_row + 1)}
        check(summary.get('检测文件数') == 3,
              '检测文件数应为 3（旧口径恒为 1），实际 %r' % summary.get('检测文件数'))
        check(summary.get('　其中：检测失败') == 1, '失败数应为 1，实际 %r' % summary.get('　其中：检测失败'))
        check(summary.get('检测目标总数') == 1, '目标总数应为 1，实际 %r' % summary.get('检测目标总数'))

        # ---------- 6) 历史批次回看：失败文件也要在 file_records 里 ----------
        r = client.get('/api/batch/BT1')
        b = r.get_json()
        check(len(b.get('file_records', [])) == 3, '历史批次台账缺失：%d' % len(b.get('file_records', [])))

        # 批次级计数：历史卡片靠它区分"全干净"和"全失败"
        st = client.get('/api/state').get_json()
        bs = (st.get('batch_stats') or {}).get('BT1') or {}
        check(bs.get('total') == 3 and bs.get('defect') == 1 and bs.get('clean') == 1
              and bs.get('failed') == 1, 'batch_stats 不对：%s' % bs)

        # ---------- 8) 单图：正常与损坏都要留下台账 ----------
        r = client.post('/api/detect/image',
                        data={'conf': '0.2', 'file': (jpeg_bytes(180), 'single.jpg')},
                        content_type='multipart/form-data')
        one = r.get_json()
        check(one.get('success') and one.get('status') == 'ok_defect',
              '单图正常路径不对：%s' % one)
        r = client.post('/api/detect/image',
                        data={'conf': '0.2', 'file': (io.BytesIO(b'bad'), 'bad_single.jpg')},
                        content_type='multipart/form-data')
        bad = r.get_json()
        check(not bad.get('success') and bad.get('status') == 'failed', '单图损坏路径不对：%s' % bad)

        # ---------- 9) 路径直读后台线程（产线主力模式）也要记台账 ----------
        src = tmp / 'src' / 'A' / '正面'
        src.mkdir(parents=True)
        from PIL import Image as _Image
        _Image.new('RGB', (64, 64), (200, 200, 200)).save(src / 'd.jpg')
        _Image.new('RGB', (64, 64), (0, 0, 0)).save(src / 'c.jpg')
        (src / 'x.jpg').write_bytes(b'bad')
        jobs = [(str(src / n), 'A/正面/' + n) for n in ('d.jpg', 'x.jpg', 'c.jpg')]
        app_module.detect_tasks['T3'] = {
            'status': 'running', 'done': 0, 'total': 3, 'defect_imgs': 0, 'defects': 0,
            'errors': 0, 'speed': 0.0, 'eta_s': -1, 'current': '', 'message': '',
            'items': [], 'stop': False, 'batch_id': 'BT3'}
        app_module._detect_worker('T3', str(tmp / 'src'), jobs, 0.2, 0, 'BT3')
        rec3 = {f['文件路径']: f for f in app_module.file_records if f.get('批次ID') == 'BT3'}
        check(len(rec3) == 3, '路径检测台账应为 3 条，实际 %d' % len(rec3))
        check(rec3.get('A/正面/d.jpg', {}).get('状态') == 'ok_defect', '亮图应 ok_defect：%s' % rec3.get('A/正面/d.jpg'))
        check(rec3.get('A/正面/c.jpg', {}).get('状态') == 'ok_clean', '黑图应 ok_clean：%s' % rec3.get('A/正面/c.jpg'))
        check(rec3.get('A/正面/x.jpg', {}).get('状态') == 'failed', '坏图应 failed：%s' % rec3.get('A/正面/x.jpg'))
        check(abs(app_module.detect_tasks['T3'].get('errors', 0) - 1) == 0,
              'errors 计数应为 1，实际 %s' % app_module.detect_tasks['T3'].get('errors'))

        # ---------- 10) 全部损坏：仍要返回 200 + 全 failed，且长度对齐 ----------
        files2 = [(io.BytesIO(b'bad'), 'a.jpg'),
                  (io.BytesIO(b'bad'), 'b.jpg')]
        r = client.post('/api/detect/batch', data={'conf': '0.3', 'batch_id': 'BT2',
                                                   'file_paths': json.dumps(['x/a.jpg', 'x/b.jpg']),
                                                   'files': files2},
                        content_type='multipart/form-data')
        resp2 = r.get_json()
        check(resp2.get('success') and resp2.get('failed') == 2 and len(resp2.get('results', [])) == 2,
              '全损坏时接口行为不对：%s' % resp2.get('failed'))
        check(all(x['status'] == 'failed' for x in resp2.get('results', [])), '全损坏时状态不对')

        # ---------- 11) 全干净批次也要能出报表（明细为空，靠台账撑住分母） ----------
        r = client.post('/api/detect/batch',
                        data={'conf': '0.3', 'batch_id': 'BT4',
                              'file_paths': json.dumps(['C/正面/ok.jpg']),
                              'files': [(jpeg_bytes(0), 'ok.jpg')]},
                        content_type='multipart/form-data')
        check(r.get_json().get('ok_clean') == 1, '全干净批次计数不对：%s' % r.get_json())
        r = client.post('/api/export', json={'batch_id': 'BT4'})
        exp2 = r.get_json()
        check(exp2.get('success'), '全干净批次导出失败（旧实现直接报"暂无检测结果"）：%s' % exp2)
        wb2 = load_workbook(results_dir / Path(exp2['download_url']).name)
        ws2 = wb2['统计汇总']
        s2 = {ws2.cell(row=i, column=1).value: ws2.cell(row=i, column=2).value
              for i in range(1, ws2.max_row + 1)}
        check(s2.get('检测文件数') == 1 and s2.get('　其中：无缺陷') == 1,
              '全干净批次汇总不对：%s' % {k: s2.get(k) for k in ('检测文件数', '　其中：无缺陷')})
        check(s2.get('检测目标总数') == 0, '全干净批次目标总数应为 0，实际 %r' % s2.get('检测目标总数'))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    if fails:
        print('\n[FAIL] %d 项' % len(fails))
        for f in fails:
            print('  -', f)
        return 1
    print('\n[PASS] Web 台账/失败态/阈值直传 全部通过')
    return 0


if __name__ == '__main__':
    sys.exit(main())
