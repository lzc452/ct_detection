# -*- coding: utf-8 -*-
"""生成《一页汇报·名词解释》Word 文档（领导版，言简意赅）"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

INK = RGBColor(0x1F, 0x32, 0x40)
ACCENT = RGBColor(0xF0, 0xA2, 0x02)
MUTED = RGBColor(0x5B, 0x6B, 0x7A)

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def set_cn(run, size=11, bold=False, color=INK):
    run.font.name = 'Times New Roman'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

def h(text, size=14, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    set_cn(p.add_run(text), size=size, bold=True)

def para(text, size=11, color=INK, bold=False, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    set_cn(p.add_run(text), size=size, bold=bold, color=color)
    return p

def kv(term, desc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_cn(p.add_run('● ' + term + '：'), bold=True)
    set_cn(p.add_run(desc))

# ===== 标题 =====
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cn(t.add_run('CT 电芯缺陷智能检测系统 · 汇报说明（名词与数字解读）'), size=16, bold=True)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cn(sub.add_run('配套一页汇报 PPT 使用 · 2026-10-19 · 戴钟科'), size=10, color=MUTED)

# ===== 一句话 =====
h('这个系统是做什么的')
para('产线 CT 会给每支电芯拍上千张"透视照片"。过去靠人眼逐张找缺陷（亮斑、暗斑、极片损伤），'
     '现在由这套系统自动找：框出缺陷位置、给出把握程度、按电芯组出报表。'
     '人只复核系统报出来的图，工作量从"看几十万张"降到"看几百张"。')

# ===== 名词解释 =====
h('名词解释（按 PPT 出现顺序）')
kv('mAP50（93.3%）', '国际通用的目标检测综合评分，兼顾"找得全"和"报得准"。'
   '可理解为期末考试成绩：60 分及格可用，80 分优秀，90 分以上属于第一梯队。')
kv('考卷', '固定的一批验证图片（196 张），每个版本模型都考同一张卷子，分数才能公平对比。')
kv('召回率', '实际存在的缺陷里，系统抓到了百分之几。召回率低 = 漏检 = 缺陷流出产线。')
kv('精确率', '系统报的缺陷里，百分之几是真的。精确率低 = 误报多 = 人工复核白费功夫。')
kv('生产口径', '按实际报警场景评分：缺陷框位置大致对（偏差几个像素也算检出）就算抓到，'
   '比国际标准宽松但更贴近产线用途。两个口径我们都测、都写明。')
kv('负样本（1,931 张）', '容易让系统误报的正常图片，专门收集起来训练，用来压误报。'
   '相当于给系统看"错题本"，这是误报大幅下降的关键手段。')
kv('数据闭环', '产线误报/漏检的图 → 人工核对 → 回流样本库 → 重新训练 → 新版本上线。'
   '已完整运转 4 轮，每轮分数提升都能说清原因。')
kv('YOLOv8n', '系统的核心算法模型，工业视觉检测的主流方案，n 代表轻量版——'
   '检测快（每张约 0.02 秒），适合产线批量扫描。')
kv('1920 高分辨率', '以接近原图的清晰度做检测。微小缺陷只有几个像素，'
   '分辨率降下来就看不见了，这是发现微小缺陷的必要代价。')

# ===== 数字解读 =====
h('四个关键数字，到底好在哪')
kv('考卷 93.3 分', '从初代 78.6 分一路提升，四代版本每一步都有明确原因'
   '（提分辨率、扩样本、清洗标注），不是碰运气的波动。')
kv('损伤召回/精确率 95.8%', '最难检的极片损伤（细到 8~14 像素宽），'
   '报 24 次 23 次是真的、24 个真缺陷抓到 23 个。上版只有 79.2%，本轮大幅提升。')
kv('样本库 4,406 张', '每一张都经人工核对、统一规范标注、可追溯来源。'
   '其中 1,931 张是专门压误报的负样本。')
kv('45 张/秒', '人工目检约 2~3 秒看一张，系统 0.02 秒检一张，效率提升超百倍。'
   '一组电芯（约 3,100 张正面图）人工需 2 小时以上，系统 1~2 分钟检完，且不受人员状态影响。')

# ===== 部署与成本 =====
h('落地方式与成本')
para('整个系统就是一个文件夹，拷贝到任何一台电脑双击即用，无需安装软件、无需联网——'
     '纯内网运行，数据不出厂。核心算法为开源方案，无版权与授权费用；'
     '硬件只需一台带普通游戏显卡的电脑（当前为 RTX 5060），无显卡也能慢速运行。')

# ===== 当前短板 =====
h('实事求是：当前的短板')
para('损伤类在更大范围抽查（649 张）中召回率约 70%，是三类中最低的。'
     '原因是产线损伤本就罕见、样本积累慢。应对措施已在执行：'
     '系统正在自动扫描 38.6 万张历史积攒图，把疑似损伤筛出来人工确认，补充进样本库后训练下一版（V5）。')

# ===== 结尾 =====
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
set_cn(p.add_run('一句话总结：'), bold=True)
set_cn(p.add_run('系统已过"能用"阶段，进入"好用"阶段——考试 93.3 分、误报较初代下降约三分之二、'
       '检测速度比人工快几个数量级，且建立了每月自动变好的数据闭环机制。'))

doc.save(r'D:\CT Detection\向上汇报\汇报说明_名词与数字解读.docx')
print('saved')
